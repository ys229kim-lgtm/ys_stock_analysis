"""
Word(.docx) 리포트 생성 스크립트 (Step 6).
생성된 마크다운과 analysis JSON을 바탕으로 python-docx로 Word 문서를 만든다.
docx 생성 실패 시 오류 로그만 남기고 종료 (md는 이미 생성된 상태).
"""
import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from statistics import mean

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def format_price(value) -> str:
    if value is None:
        return "N/A"
    return f"{int(value):,}원"


def set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_cover_page(doc: Document, company: str, date_fmt: str, report_count: int, period_str: str) -> None:
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{company} 증권사 리포트 분석")
    run.font.size = Pt(24)
    run.font.bold = True

    doc.add_paragraph()
    for label, value in [
        ("분석일자", date_fmt),
        ("수집 리포트 수", f"{report_count}건"),
        ("수집 기간", period_str),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}: {value}")
        run.font.size = Pt(14)

    doc.add_page_break()


def add_opinion_table(doc: Document, reports: list[dict]) -> None:
    counts = defaultdict(int)
    for r in reports:
        op = r.get("investment_opinion", "명시없음") or "명시없음"
        counts[op] += 1
    total = sum(counts.values())

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["투자의견", "건수", "비율"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].font.bold = True
        set_cell_bg(hdr[i], "D0E4F7")

    for op in ["Buy", "Hold", "Sell", "명시없음"]:
        n = counts.get(op, 0)
        pct = f"{n/total*100:.0f}%" if total > 0 else "0%"
        row = table.add_row().cells
        row[0].text = op
        row[1].text = str(n)
        row[2].text = pct


def add_scenario_box(doc: Document, title: str, content: str, color: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"▌ {title}")
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(13)

    for line in content.strip().split("\n"):
        doc.add_paragraph(line, style="List Bullet")


def parse_synthesis_sections(synthesis_text: str) -> dict[str, str]:
    sections: dict[str, str] = {}
    current_key = None
    current_lines: list[str] = []

    heading_re = re.compile(r"^#{1,3}\s+(.+)")

    for line in synthesis_text.split("\n"):
        m = heading_re.match(line)
        if m:
            if current_key:
                sections[current_key] = "\n".join(current_lines).strip()
            current_key = m.group(1).strip()
            current_lines = []
        else:
            if current_key:
                current_lines.append(line)

    if current_key:
        sections[current_key] = "\n".join(current_lines).strip()

    return sections


def build_docx(
    company: str,
    date: str,
    meta: dict,
    analysis: dict,
    synthesis_text: str,
) -> Document:
    doc = Document()

    reports = [
        r for r in analysis.get("reports", [])
        if not r.get("irrelevant") and not r.get("parse_failed")
    ]
    date_fmt = f"{date[:4]}-{date[4:6]}-{date[6:8]}"
    period_from = meta.get("period", {}).get("from", "")
    period_to = meta.get("period", {}).get("to", date_fmt)
    period_str = f"{period_from} ~ {period_to}" if period_from else f"~ {period_to}"

    add_cover_page(doc, company, date_fmt, meta.get("after_dedup", 0), period_str)

    sections = parse_synthesis_sections(synthesis_text)

    doc.add_heading("1. 수집 개요", level=1)
    p = doc.add_paragraph()
    p.add_run(f"총 수집: {meta.get('total_collected', 0)}건  /  "
              f"중복 제거 후: {meta.get('after_dedup', 0)}건  /  "
              f"유효 분석: {len(reports)}건")

    by_channel = meta.get("by_channel", {})
    if by_channel:
        doc.add_paragraph("채널별 수집:", style="Heading 3")
        ch_table = doc.add_table(rows=1, cols=2)
        ch_table.style = "Table Grid"
        ch_table.rows[0].cells[0].text = "채널"
        ch_table.rows[0].cells[1].text = "건수"
        for ch, n in sorted(by_channel.items(), key=lambda x: -x[1]):
            row = ch_table.add_row().cells
            row[0].text = ch
            row[1].text = str(n)

    doc.add_heading("2. 투자의견 분포", level=1)
    add_opinion_table(doc, reports)
    prices = [r["target_price"] for r in reports if r.get("target_price") is not None]
    if prices:
        doc.add_paragraph(
            f"목표주가 범위: 최소 {format_price(min(prices))} / "
            f"평균 {format_price(int(mean(prices)))} / "
            f"최대 {format_price(max(prices))}"
        )

    for heading in ["3. 긍정 논거 통합", "4. 부정 논거 통합", "5. 실적 전망 컨센서스"]:
        doc.add_heading(heading, level=1)
        matched = next((v for k, v in sections.items() if heading[3:].strip() in k), "")
        doc.add_paragraph(matched or "(데이터 없음)")

    doc.add_heading("6. Best Case 시나리오", level=1)
    best_content = next((v for k, v in sections.items() if "Best" in k or "best" in k.lower()), "")
    add_scenario_box(doc, "Best Case", best_content or "(데이터 없음)", "1F6AA5")

    doc.add_heading("7. Worst Case 시나리오", level=1)
    worst_content = next((v for k, v in sections.items() if "Worst" in k or "worst" in k.lower()), "")
    add_scenario_box(doc, "Worst Case", worst_content or "(데이터 없음)", "C0392B")

    doc.add_heading("8. 종합 시각 요약", level=1)
    summary_content = next(
        (v for k, v in sections.items() if "종합" in k or "요약" in k), ""
    )
    doc.add_paragraph(summary_content or "(데이터 없음)")

    doc.add_page_break()
    doc.add_heading("부록: 리포트별 분석 상세", level=1)
    app_table = doc.add_table(rows=1, cols=5)
    app_table.style = "Table Grid"
    headers = ["증권사", "날짜", "투자의견", "목표주가", "채널"]
    for i, h in enumerate(headers):
        app_table.rows[0].cells[i].text = h
        set_cell_bg(app_table.rows[0].cells[i], "D0E4F7")

    for r in sorted(reports, key=lambda x: x.get("report_date", ""), reverse=True):
        row = app_table.add_row().cells
        row[0].text = r.get("securities_firm") or "미상"
        row[1].text = r.get("report_date", "")
        row[2].text = r.get("investment_opinion", "명시없음")
        row[3].text = format_price(r.get("target_price"))
        row[4].text = r.get("channel", "")

    return doc


def main():
    if not HAS_DOCX:
        print("[generate_docx] ERROR: python-docx 미설치. `pip install python-docx` 실행 필요.")
        raise SystemExit(1)

    parser = argparse.ArgumentParser(description="Word 문서 생성")
    parser.add_argument("--analysis-json", required=True)
    parser.add_argument("--synthesis-text", required=True)
    parser.add_argument("--meta-json", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--company", required=True)
    parser.add_argument("--date", required=True)
    args = parser.parse_args()

    analysis = load_json(Path(args.analysis_json))
    meta = load_json(Path(args.meta_json))
    synthesis_text = Path(args.synthesis_text).read_text(encoding="utf-8")

    try:
        doc = build_docx(args.company, args.date, meta, analysis, synthesis_text)
        output_dir = Path(args.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        docx_path = output_dir / f"report_{args.date}_{args.company}.docx"
        doc.save(str(docx_path))
        print(f"[generate_docx] Word 문서 저장: {docx_path}")
    except Exception as e:
        print(f"[generate_docx] ERROR: docx 생성 실패 — {e}")
        print("[generate_docx] md 파일로 계속 진행합니다.")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
