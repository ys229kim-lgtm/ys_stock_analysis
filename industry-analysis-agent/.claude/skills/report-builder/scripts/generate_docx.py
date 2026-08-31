"""
Step 5: analysis JSON → Word(.docx) 리포트 생성 (python-docx)
Usage: python generate_docx.py --json_paths ./output/analysis_20260608_삼성전자_반도체.json
                                --out_path ./output/report_20260608_삼성전자.docx
복수 산업: --json_paths path1.json path2.json (순서대로 섹션 이어서 작성)
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# 시나리오 셀 배경색
COLOR_POSITIVE = "DDEEFF"
COLOR_BASE     = "F0F0F0"
COLOR_NEGATIVE = "FFDDDD"


def set_cell_background(cell, hex_color: str):
    """셀 배경색 설정."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_heading(doc, text: str, level: int):
    doc.add_heading(text, level=level)


def add_value_chain_table(doc, value_chain: list):
    if not value_chain:
        doc.add_paragraph("밸류체인 데이터 없음")
        return

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "단계"
    hdr[1].text = "주요 플레이어"

    for item in value_chain:
        row = table.add_row().cells
        row[0].text = item.get("stage", "")
        row[1].text = ", ".join(item.get("players", []))


def add_scenario_table(doc, horizon_data: dict):
    scenarios = [
        ("positive", "🟢 Positive", COLOR_POSITIVE),
        ("base",     "⚪ Base Case", COLOR_BASE),
        ("negative", "🔴 Negative",  COLOR_NEGATIVE),
    ]

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["시나리오", "확률", "핵심 조건", "시장 성장 전망", "대상 기업 영향"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for key, label, bg_color in scenarios:
        s = horizon_data.get(key, {})
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(s.get("probability", "N/A"))
        row[2].text = ", ".join(s.get("conditions", []))
        row[3].text = str(s.get("market_growth_est", "N/A"))
        row[4].text = str(s.get("impact_on_target", "N/A"))
        for cell in row:
            set_cell_background(cell, bg_color)


def add_industry_section(doc, data: dict, is_first: bool = True):
    company = data.get("company", "")
    industry = data.get("industry", "")
    date = data.get("analysis_date", "")

    if not is_first:
        doc.add_page_break()

    add_heading(doc, f"{industry} 산업 분석", level=1)

    overview = data.get("overview", {})
    market_size = data.get("market_size", {})
    competition = data.get("competition", {})
    supply_demand = data.get("supply_demand", {})
    global_trends = data.get("global_trends", [])
    regulation = data.get("regulation", {})
    scenarios = data.get("scenarios", {})
    conclusion = data.get("conclusion", {})

    # 1. 산업 개요
    add_heading(doc, "1. 산업 개요 및 밸류체인", level=2)
    doc.add_paragraph(overview.get("industry_definition", "정의 없음"))
    doc.add_paragraph(f"대상 기업 위치: {overview.get('target_company_position', '미기재')}")
    add_value_chain_table(doc, overview.get("value_chain", []))

    # 2. 시장 규모
    add_heading(doc, "2. 시장 규모 및 성장률 전망", level=2)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    headers = [["구분", "규모", "비고"],
               ["현재", str(market_size.get("current_global", {}).get("value", "")), f"출처: {market_size.get('source','')}"],
               ["단기 전망(1년)", str(market_size.get("forecast_1y", {}).get("value", "")), f"CAGR: {market_size.get('cagr_1y','')}"],
               ["중기 전망(3년)", str(market_size.get("forecast_3y", {}).get("value", "")), f"CAGR: {market_size.get('cagr_3y','')}"]]
    for i, row_data in enumerate(headers):
        for j, val in enumerate(row_data):
            table.rows[i].cells[j].text = val

    # 3. 경쟁 구도
    add_heading(doc, "3. 주요 플레이어 경쟁 구도", level=2)
    top_players = competition.get("top_players", [])
    if top_players:
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for cell, h in zip(tbl.rows[0].cells, ["기업", "점유율", "지역"]):
            cell.text = h
        for p in top_players:
            row = tbl.add_row().cells
            row[0].text = p.get("company", "")
            row[1].text = p.get("market_share", "")
            row[2].text = p.get("region", "")
    doc.add_paragraph(f"대상 기업 포지션: {competition.get('target_position', '미기재')}")
    doc.add_paragraph(f"경쟁 강도: {competition.get('competitive_intensity', '미기재')}")

    # 4. 수급 사이클
    add_heading(doc, "4. 수요·공급 사이클 분석", level=2)
    doc.add_paragraph(f"현재 사이클 위치: {supply_demand.get('current_cycle', '미기재')}")
    doc.add_paragraph(f"재고 수준: {supply_demand.get('inventory_level', '미기재')}")
    doc.add_paragraph(f"가격 추이: {supply_demand.get('price_trend', '미기재')}")

    # 5. 글로벌 트렌드
    add_heading(doc, "5. 글로벌 트렌드 및 기술 변화", level=2)
    for trend in global_trends:
        doc.add_paragraph(f"• {trend}")

    # 6. 규제 환경
    add_heading(doc, "6. 규제 환경", level=2)
    doc.add_paragraph("국내:")
    for r in regulation.get("domestic", []):
        doc.add_paragraph(f"• {r}")
    doc.add_paragraph("해외:")
    for r in regulation.get("overseas", []):
        doc.add_paragraph(f"• {r}")

    # 7. 단기 전망
    add_heading(doc, "7. 단기(1년) 확률적 방향성 전망", level=2)
    add_scenario_table(doc, scenarios.get("short_term", {}))

    # 8. 중기 전망
    add_heading(doc, "8. 중기(1~3년) 확률적 방향성 전망", level=2)
    add_scenario_table(doc, scenarios.get("mid_term", {}))
    direction = scenarios.get("direction_consistency", "")
    if direction:
        doc.add_paragraph(f"단기·중기 방향성: {direction}")

    # 9. 종합 평가
    add_heading(doc, "9. 종합 평가 및 대상 기업 시사점", level=2)
    if conclusion:
        doc.add_paragraph(f"① 시장 순풍 vs 역풍\n{conclusion.get('tailwind_headwind','미작성')}")
        doc.add_paragraph(f"② 포지셔닝 개선 여지\n{conclusion.get('positioning_improvement','미작성')}")
        doc.add_paragraph(f"③ 할인율·멀티플 연결\n{conclusion.get('valuation_linkage','미작성')}")


def build_cover_page(doc, company: str, industries: list[str], date: str, sources: dict):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{company}")
    run.bold = True
    run.font.size = Pt(28)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(" · ".join(industries) + " 산업 분석")
    run2.font.size = Pt(18)

    doc.add_paragraph(f"분석일자: {date}").alignment = WD_ALIGN_PARAGRAPH.CENTER

    src = sources
    src_text = (
        f"텔레그램 리포트 {src.get('telegram_reports',0)}건 · "
        f"웹 기사 {src.get('web_articles',0)}건 · "
        f"시장 데이터 {src.get('market_tickers',0)}개 기업"
    )
    doc.add_paragraph(src_text).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def main():
    if not DOCX_AVAILABLE:
        print("[generate_docx] python-docx 미설치. pip install python-docx 실행 후 재시도.", file=sys.stderr)
        sys.exit(1)

    parser = argparse.ArgumentParser()
    parser.add_argument("--json_paths", nargs="+", required=True, help="analysis JSON 파일 경로 (복수 가능)")
    parser.add_argument("--out_path", required=True, help="출력 docx 경로")
    args = parser.parse_args()

    all_data = []
    for p in args.json_paths:
        path = Path(p)
        if not path.exists():
            print(f"[generate_docx] 파일 없음: {path}", file=sys.stderr)
            continue
        all_data.append(json.loads(path.read_text(encoding="utf-8")))

    if not all_data:
        print("[generate_docx] 유효한 JSON 파일 없음", file=sys.stderr)
        sys.exit(1)

    doc = Document()

    # 표지
    company = all_data[0].get("company", "")
    industries = [d.get("industry", "") for d in all_data]
    date = all_data[0].get("analysis_date", datetime.now().strftime("%Y-%m-%d"))
    combined_sources = {
        "telegram_reports": sum(d.get("data_sources", {}).get("telegram_reports", 0) for d in all_data),
        "web_articles": sum(d.get("data_sources", {}).get("web_articles", 0) for d in all_data),
        "market_tickers": sum(d.get("data_sources", {}).get("market_tickers", 0) for d in all_data),
    }
    build_cover_page(doc, company, industries, date, combined_sources)

    # 각 산업 섹션
    for i, data in enumerate(all_data):
        add_industry_section(doc, data, is_first=(i == 0))

    # 부록
    all_refs = []
    for d in all_data:
        all_refs.extend(d.get("references", []))
    if all_refs:
        doc.add_page_break()
        add_heading(doc, "부록: 참고 자료", level=1)
        for ref in all_refs:
            doc.add_paragraph(f"• {ref}")

    out_path = Path(args.out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"[generate_docx] Word 리포트 생성: {out_path}")


if __name__ == "__main__":
    main()
