#!/usr/bin/env python3
"""
Step 7-3: Word 문서 생성 (python-docx)
Markdown 리포트를 읽어 .docx로 변환한다.
표지·섹션 스타일·표 서식 적용.
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("[ERROR] python-docx 미설치. 아래 명령으로 설치하세요:", file=sys.stderr)
    print("  pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ─────────────────────────────────────────────
# 색상 상수
# ─────────────────────────────────────────────
COLOR_H1 = RGBColor(0x1F, 0x38, 0x64)   # 진한 파란색
COLOR_H2 = RGBColor(0x2F, 0x54, 0x96)   # 중간 파란색
COLOR_H3 = RGBColor(0x00, 0x00, 0x00)   # 검정
COLOR_TABLE_HEADER_BG = "1F3864"         # 헤더 배경 (OOXML 형식)
COLOR_TABLE_ROW_ALT = "F2F2F2"           # 번갈아 행 배경
COLOR_POSITIVE = RGBColor(0x00, 0x70, 0x00)  # 양수 녹색
COLOR_NEGATIVE = RGBColor(0xC0, 0x00, 0x00)  # 음수 빨강


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_cover_page(doc: Document, company: str, date: str, summary: dict):
    date_fmt = f"{date[:4]}-{date[4:6]}-{date[6:8]}" if len(date) == 8 else date
    current_price = summary.get("current_price", "N/A")
    fair_value = summary.get("weighted_avg_fair_value", "N/A")
    upside = summary.get("upside_pct")
    upside_str = f"{upside:+.1f}%" if isinstance(upside, (int, float)) else "N/A"
    fvr = summary.get("fair_value_range") or {}

    # 표지 제목
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{company}")
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = COLOR_H1

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("기업분석 종합 리포트")
    run2.font.size = Pt(20)
    run2.font.color.rgb = COLOR_H2

    doc.add_paragraph()

    # 핵심 수치 표
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows_data = [
        ("분석일자", date_fmt),
        ("현재 주가", f"{current_price:,}원" if isinstance(current_price, (int, float)) else str(current_price)),
        ("가중 평균 적정가", f"{fair_value:,}원" if isinstance(fair_value, (int, float)) else str(fair_value)),
        ("적정가 범위", f"{fvr.get('low', 'N/A')}원 ~ {fvr.get('high', 'N/A')}원"),
        ("Upside/Downside", upside_str),
    ]
    for i, (key, val) in enumerate(rows_data):
        row = table.rows[i]
        row.cells[0].text = key
        row.cells[1].text = val
        set_cell_bg(row.cells[0], "2F5496")
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        if i % 2 == 1:
            set_cell_bg(row.cells[1], COLOR_TABLE_ROW_ALT)

    doc.add_page_break()


def apply_heading_style(paragraph, level: int):
    for run in paragraph.runs:
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = COLOR_H1
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = COLOR_H2
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = COLOR_H3


def parse_md_table(lines: list) -> list[list[str]]:
    """마크다운 표 파싱 → [[cell, ...], ...]"""
    rows = []
    for line in lines:
        line = line.strip()
        if not line or set(line.replace("|", "").replace("-", "").replace(" ", "")) == set():
            continue
        if re.match(r"^\|[-| :]+\|$", line):
            continue
        cells = [c.strip() for c in line.split("|")]
        cells = [c for c in cells if c != ""]
        if cells:
            rows.append(cells)
    return rows


def add_md_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    max_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"

    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            if c_idx >= max_cols:
                break
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = cell_text
            run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(cell_text)

            if r_idx == 0:
                # 헤더 행
                set_cell_bg(cell, COLOR_TABLE_HEADER_BG)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True
            elif r_idx % 2 == 0:
                set_cell_bg(cell, COLOR_TABLE_ROW_ALT)

            # upside% 컬럼 색상 처리
            if "%" in cell_text and r_idx > 0:
                try:
                    num = float(re.sub(r"[^0-9.\-+]", "", cell_text))
                    if num > 0:
                        run.font.color.rgb = COLOR_POSITIVE
                    elif num < 0:
                        run.font.color.rgb = COLOR_NEGATIVE
                except ValueError:
                    pass


def process_markdown(doc: Document, md_content: str):
    """마크다운 텍스트를 파싱하여 docx에 추가."""
    lines = md_content.split("\n")
    i = 0
    table_buffer = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # 표 처리
        if line.strip().startswith("|"):
            in_table = True
            table_buffer.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            rows = parse_md_table(table_buffer)
            add_md_table(doc, rows)
            table_buffer = []

        # 구분선
        if line.strip() in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "2F5496")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # 헤딩
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            run = p.add_run(text)
            apply_heading_style(p, level)
            i += 1
            continue

        # 빈 줄
        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue

        # 일반 단락
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        # 인라인 볼드 처리
        parts = re.split(r"(\*\*[^*]+\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            else:
                p.add_run(part)

        i += 1

    # 마지막에 표가 남아있으면 처리
    if table_buffer:
        rows = parse_md_table(table_buffer)
        add_md_table(doc, rows)


def set_page_margins(doc: Document):
    from docx.oxml.ns import qn
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def main():
    parser = argparse.ArgumentParser(description="Word 문서 생성")
    parser.add_argument("--company", required=True)
    parser.add_argument("--date", required=True)
    parser.add_argument("--final-path", required=True)
    args = parser.parse_args()

    final_dir = Path(args.final_path)
    md_path = final_dir / f"report_{args.date}_{args.company}.md"
    docx_path = final_dir / f"report_{args.date}_{args.company}.docx"

    if not md_path.exists():
        print(f"[ERROR] Markdown 파일 없음: {md_path}", file=sys.stderr)
        sys.exit(1)

    with open(md_path, encoding="utf-8") as f:
        md_content = f.read()

    # integrated JSON에서 summary 추출
    summary = {}
    integrated_path = final_dir / f"integrated_{args.date}_{args.company}.json"
    if integrated_path.exists():
        try:
            with open(integrated_path, encoding="utf-8") as f:
                integrated = json.load(f)
            summary = integrated.get("summary", {})
        except Exception as e:
            print(f"[WARN] integrated JSON 로딩 실패: {e}", file=sys.stderr)

    # Word 문서 생성
    doc = Document()
    set_page_margins(doc)

    # 기본 폰트 설정
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    # 표지 삽입
    add_cover_page(doc, args.company, args.date, summary)

    # 마크다운 본문 처리
    # 표지 부분(첫 --- 블록) 스킵하고 본문만 처리
    body_start = md_content.find("\n---\n", md_content.find("기업분석 종합 리포트"))
    if body_start >= 0:
        body_content = md_content[body_start + 5:]
    else:
        body_content = md_content

    process_markdown(doc, body_content)

    doc.save(docx_path)
    size_kb = docx_path.stat().st_size // 1024
    print(f"[OK] Word 문서 저장: {docx_path} ({size_kb}KB)", file=sys.stderr)
    print(json.dumps({"output": str(docx_path), "size_kb": size_kb}, ensure_ascii=False))


if __name__ == "__main__":
    main()
