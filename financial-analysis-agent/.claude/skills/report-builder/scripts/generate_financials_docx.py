#!/usr/bin/env python3
"""
generate_financials_docx.py — 재무 수치 테이블 Word 문서 생성 (python-docx)

Usage:
    python generate_financials_docx.py \
        --financials financials.json \
        --output financials.docx \
        [--competitors comp1.json comp2.json]
"""

import argparse
import json
from pathlib import Path
from datetime import datetime


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--financials",  required=True)
    parser.add_argument("--output",      required=True)
    parser.add_argument("--competitors", nargs="*", default=[])
    args = parser.parse_args()

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("[ERROR] python-docx 미설치 — pip install python-docx", flush=True)
        raise SystemExit(1)

    with open(args.financials, encoding="utf-8") as f:
        fin = json.load(f)

    comp_data = []
    for cp in args.competitors:
        with open(cp, encoding="utf-8") as f:
            comp_data.append(json.load(f))

    doc = Document()
    doc.core_properties.author = "Financial Analysis Agent"

    company = fin["company"]["name"]
    rtype   = fin["company"].get("report_type", "연결")
    dr      = fin.get("data_range", {})
    date_str= datetime.now().strftime("%Y년 %m월 %d일")

    # Cover
    t = doc.add_heading(company, 0); t.alignment = 1
    for txt in [
        f"재무실적 데이터  |  분석일자: {date_str}",
        f"분석 기간: {dr.get('start_year','')}–{dr.get('end_year','')}년  |  기준: {rtype}",
    ]:
        p = doc.add_paragraph(txt); p.alignment = 1
    doc.add_page_break()

    annual = sorted(fin.get("annual", []), key=lambda a: a["year"])
    years  = [a["year"] for a in annual]

    def gv(a, *keys):
        v = a
        for k in keys:
            v = v.get(k) if isinstance(v, dict) else None
        return round(v, 2) if v is not None else None

    def add_table(title, rows_def):
        doc.add_heading(title, level=1)
        tbl = doc.add_table(rows=1 + len(rows_def), cols=1 + len(years))
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text = "지표"
        for i, y in enumerate(years):
            hdr[i+1].text = str(y)
        for ri, (label, vals) in enumerate(rows_def):
            row = tbl.rows[ri+1].cells
            row[0].text = label
            for ci, v in enumerate(vals):
                row[ci+1].text = str(v) if v is not None else "N/A"
        doc.add_paragraph()

    add_table("1. 포괄손익계산서 (단위: 조원)", [
        ("매출액",        [gv(a,"income_statement","revenue")          for a in annual]),
        ("영업이익",      [gv(a,"income_statement","operating_profit")  for a in annual]),
        ("당기순이익",    [gv(a,"income_statement","net_income")        for a in annual]),
        ("영업이익률(%)", [gv(a,"derived_metrics","operating_margin")   for a in annual]),
        ("순이익률(%)",   [gv(a,"derived_metrics","net_margin")         for a in annual]),
    ])

    add_table("2. 재무상태표 (단위: 조원)", [
        ("자산총계",  [gv(a,"balance_sheet","total_assets")       for a in annual]),
        ("자본총계",  [gv(a,"balance_sheet","total_equity")       for a in annual]),
        ("부채총계",  [gv(a,"balance_sheet","total_debt")         for a in annual]),
        ("유동자산",  [gv(a,"balance_sheet","current_assets")     for a in annual]),
        ("유동부채",  [gv(a,"balance_sheet","current_liabilities") for a in annual]),
    ])

    add_table("3. 현금흐름표 (단위: 조원)", [
        ("영업CF",  [gv(a,"cash_flow","operating_cf")  for a in annual]),
        ("투자CF",  [gv(a,"cash_flow","investing_cf")  for a in annual]),
        ("재무CF",  [gv(a,"cash_flow","financing_cf")  for a in annual]),
        ("CAPEX",   [gv(a,"cash_flow","capex")          for a in annual]),
        ("FCF",     [gv(a,"derived_metrics","fcf")      for a in annual]),
    ])

    add_table("4. 파생 재무지표", [
        ("ROE(%)",          [gv(a,"derived_metrics","roe")               for a in annual]),
        ("ROA(%)",          [gv(a,"derived_metrics","roa")               for a in annual]),
        ("부채비율(%)",     [gv(a,"derived_metrics","debt_ratio")        for a in annual]),
        ("유동비율(%)",     [gv(a,"derived_metrics","current_ratio")     for a in annual]),
        ("이자보상배율(x)", [gv(a,"derived_metrics","interest_coverage") for a in annual]),
        ("FCF 마진(%)",     [gv(a,"derived_metrics","fcf_margin")        for a in annual]),
    ])

    # 경쟁사 비교
    if comp_data:
        doc.add_page_break()
        doc.add_heading("5. 경쟁사 비교 (최근 5개년)", level=1)
        all_co  = [fin] + comp_data
        all_y5  = sorted({a["year"] for d in all_co for a in d["annual"]})[-5:]

        for metric, sec, field in [
            ("영업이익률(%)",  "derived_metrics",  "operating_margin"),
            ("ROE(%)",         "derived_metrics",  "roe"),
            ("부채비율(%)",    "derived_metrics",  "debt_ratio"),
            ("FCF(조원)",      "derived_metrics",  "fcf"),
        ]:
            doc.add_heading(metric, level=2)
            tbl = doc.add_table(rows=1+len(all_co), cols=1+len(all_y5))
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "기업"
            for i, y in enumerate(all_y5): hdr[i+1].text = str(y)
            for ri, co in enumerate(all_co):
                row = tbl.rows[ri+1].cells
                row[0].text = co["company"]["name"]
                ym = {a["year"]: a for a in co.get("annual", [])}
                for ci, y in enumerate(all_y5):
                    a = ym.get(y, {})
                    v = a.get(sec, {}).get(field) if a else None
                    row[ci+1].text = str(round(v, 2)) if v is not None else "N/A"
            doc.add_paragraph()

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[OK] financials.docx 저장: {out}")


if __name__ == "__main__":
    main()
